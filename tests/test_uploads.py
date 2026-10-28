"""Upload limits + hostile-file tests. Every 'attack' file here must be refused with a clear reason."""
import base64
import html
import io
import json
import os
import re
import zipfile
from datetime import datetime

import pytest
import segno
from openpyxl import Workbook

MB = 1024 * 1024


# ------------------------------------------------------------------ helpers
def xlsx(rows=None, sheets=None) -> bytes:
    wb = Workbook()
    if sheets:
        wb.remove(wb.active)
        for name, data in sheets.items():
            ws = wb.create_sheet(name)
            for r in data:
                ws.append(r)
    else:
        for r in rows or [["name", "qty", "when"], ["apple", 3, datetime(2026, 1, 2, 3, 4, 5)], ["pear", 4.5, None]]:
            wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def with_parts(data: bytes, extra: dict | None = None, replace: dict | None = None) -> bytes:
    """Rebuild an xlsx with extra/replaced parts (to craft malicious workbooks)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            body = (replace or {}).get(info.filename, src.read(info.filename))
            z.writestr(info.filename, body)
        for name, body in (extra or {}).items():
            z.writestr(name, body)
    return out.getvalue()


def png(side=64, noise=False) -> bytes:
    from PIL import Image
    img = Image.frombytes("RGB", (side, side), os.urandom(side * side * 3)) if noise else Image.new("RGB", (side, side), "teal")
    buf = io.BytesIO()
    img.save(buf, "PNG", compress_level=0 if noise else 6)
    return buf.getvalue()


def text(resp) -> str:
    return html.unescape(re.sub(r"<[^>]+>", "", resp.get_data(as_text=True)))


def up(run, slug, data: bytes, name="upload.bin", **fields):
    fields.setdefault("action", None)
    if fields["action"] is None:
        del fields["action"]
    return run(slug, file=(io.BytesIO(data), name), **fields)


# ------------------------------------------------------------------ configuration
def test_default_limits_are_the_requested_ones(app):
    lim = app.config["UPLOAD_LIMITS"]
    assert lim["excel-json"] == 1 * MB
    assert lim["base64-image"] == 5 * MB
    assert lim["certificate"] == 2 * MB
    assert lim["qr-code"] == 2 * MB
    assert app.config["MAX_CONTENT_LENGTH"] >= 5 * MB


def test_limits_can_be_overridden_by_env_and_are_capped(monkeypatch):
    from app.config import _limit_mb
    monkeypatch.setenv("UPLOAD_MB_EXCEL_JSON", "3")
    assert _limit_mb("UPLOAD_MB_EXCEL_JSON", 1) == 3 * MB
    monkeypatch.setenv("UPLOAD_MB_EXCEL_JSON", "9999")
    assert _limit_mb("UPLOAD_MB_EXCEL_JSON", 1) == 20 * MB           # hard ceiling
    monkeypatch.setenv("UPLOAD_MB_EXCEL_JSON", "junk")
    assert _limit_mb("UPLOAD_MB_EXCEL_JSON", 1) == 1 * MB            # bad value -> default


def test_uploads_never_touch_disk(app):
    from app import SafeRequest
    assert app.request_class is SafeRequest
    assert isinstance(SafeRequest({})._get_file_stream(1, "x"), io.BytesIO)


def test_page_shows_limit(client):
    assert "Max file size: 1 MB" in client.get("/tool/excel-json").get_data(as_text=True)
    assert "Max file size: 5 MB" in client.get("/tool/base64-image").get_data(as_text=True)
    assert "Max file size: 2 MB" in client.get("/tool/certificate").get_data(as_text=True)
    assert "Max file size: 2 MB" in client.get("/tool/qr-code").get_data(as_text=True)
    assert 'data-max-bytes="1048576"' in client.get("/tool/excel-json").get_data(as_text=True)


# ------------------------------------------------------------------ size limits per tool
def test_excel_over_1mb_is_refused(run):
    big = b"PK\x03\x04" + b"\0" * (1 * MB + 10)
    assert "larger than the 1 MB limit" in text(up(run, "excel-json", big, "a.xlsx"))


def test_excel_gross_oversize_refused_early_with_413(client):
    r = client.post("/tool/excel-json/run", data=b"x" * (3 * MB), content_type="application/octet-stream")
    assert r.status_code == 413


def test_valid_xlsx_just_under_limit_is_accepted(run):
    data = xlsx([["id", "v"]] + [[i, "x" * 20] for i in range(500)])
    assert len(data) < 1 * MB
    r = up(run, "excel-json", data, "ok.xlsx", header="on")
    assert "Rejected" not in text(r) and "Security check passed" in text(r)


def test_base64_image_allows_4mb_but_not_over_5mb(run):
    ok = png(1200, noise=True)                    # ~4.3 MB
    assert 4 * MB < len(ok) < 5 * MB
    r = up(run, "base64-image", ok, "big.png", action="encode")
    assert r.status_code == 200 and "larger than" not in text(r) and "Rejected" not in text(r)
    too_big = png(1400, noise=True)               # ~5.9 MB (under the 6 MB request ceiling)
    assert len(too_big) > 5 * MB
    assert "larger than the 5 MB limit" in text(up(run, "base64-image", too_big, "huge.png", action="encode"))


def test_certificate_over_2mb_refused(run):
    blob = b"-----BEGIN CERTIFICATE-----\n" + b"A" * (2 * MB + 5)
    assert "larger than the 2 MB limit" in text(up(run, "certificate", blob, "c.pem", action="decode"))


def test_qr_over_2mb_refused(run):
    blob = png(900, noise=True)
    assert len(blob) > 2 * MB
    assert "larger than the 2 MB limit" in text(up(run, "qr-code", blob, "q.png", action="decode"))


# ------------------------------------------------------------------ Excel to JSON behaviour
def test_excel_converts_rows_to_objects(run):
    r = up(run, "excel-json", xlsx(), "data.xlsx", header="on", indent="2")
    rows = _out_json(r)["Sheet"]
    assert rows[0] == {"name": "apple", "qty": 3, "when": "2026-01-02T03:04:05"}
    assert rows[1]["qty"] == 4.5 and rows[1].get("when") is None
    assert "Security check passed" in text(r)


def _out_json(resp):
    body = resp.get_data(as_text=True).split('<pre id="out" class="hl">')[1].split("</pre>")[0]
    return json.loads(html.unescape(re.sub(r"<[^>]+>", "", body)))


def test_excel_single_sheet_array_and_no_header(run):
    data = xlsx(sheets={"A": [["x", "y"], [1, 2]], "B": [["z"], [9]]})
    assert _out_json(up(run, "excel-json", data, "d.xlsx", header="on", sheet="B")) == [{"z": 9}]
    assert _out_json(up(run, "excel-json", data, "d.xlsx", sheet="2")) == [["z"], [9]]
    both = _out_json(up(run, "excel-json", data, "d.xlsx", header="on"))
    assert set(both) == {"A", "B"}
    assert "Sheet not found" in text(up(run, "excel-json", data, "d.xlsx", sheet="nope"))


def test_excel_duplicate_and_blank_headers(run):
    rows = _out_json(up(run, "excel-json", xlsx(sheets={"S": [["a", "a", None], [1, 2, 3]]}), "d.xlsx", header="on", sheet="S"))
    assert rows == [{"a": 1, "a_2": 2, "column_3": 3}]


def test_excel_formulas_are_not_evaluated(run):
    wb = Workbook(); ws = wb.active
    ws.append(["f"]); ws.append(["=1+1"]); ws.append(['=HYPERLINK("http://evil","x")'])
    buf = io.BytesIO(); wb.save(buf)
    r = up(run, "excel-json", buf.getvalue(), "f.xlsx", header="on", sheet="1")
    assert _out_json(r) == []                       # formula cells have no cached value -> skipped, never evaluated
    page = r.get_data(as_text=True)
    assert "HYPERLINK" not in page and "evil" not in page and "1+1" not in page


def test_excel_output_is_escaped(run):
    data = xlsx(sheets={"<img src=x onerror=alert(1)>": [["h"], ["<script>alert(1)</script>"]]})
    page = up(run, "excel-json", data, "x.xlsx", header="on").get_data(as_text=True)
    assert "<script" not in page and "<img" not in page and "&lt;script&gt;" in page


def test_excel_row_cap(run):
    wb = Workbook(write_only=True); ws = wb.create_sheet()
    for i in range(20_005):
        ws.append([i])
    buf = io.BytesIO(); wb.save(buf)
    assert len(buf.getvalue()) < 1 * MB
    assert "more than 20,000 rows" in text(up(run, "excel-json", buf.getvalue(), "r.xlsx", header="on"))


def test_client_filename_and_type_are_ignored(run):
    assert "Rejected" not in text(up(run, "excel-json", xlsx(), "totally-an.exe", header="on"))   # valid content wins
    assert "Windows executable" in text(up(run, "excel-json", b"MZ\x90\x00" + b"\0" * 200, "report.xlsx"))


# ------------------------------------------------------------------ hostile Excel files
@pytest.mark.parametrize("label,blob,expect", [
    ("renamed exe", b"MZ\x90\x00\x03" + b"\0" * 300, "Windows executable"),
    ("renamed ELF", b"\x7fELF" + b"\0" * 300, "Linux executable"),
    ("shell script", b"#!/bin/sh\nrm -rf /\n", "script"),
    ("legacy xls / OLE", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 300, "legacy Office"),
    ("php / html", b"<?php system($_GET['c']); ?>", "script or markup"),
    ("pdf", b"%PDF-1.7\n" + b"0" * 100, "PDF"),
    ("random text", b"just some text, not a workbook", "not an .xlsx"),
    ("empty zip", b"PK\x05\x06" + b"\0" * 18, "not an .xlsx"),
])
def test_hostile_excel_uploads_rejected(run, label, blob, expect):
    t = text(up(run, "excel-json", blob, "evil.xlsx"))
    assert "Rejected" in t and expect in t, (label, t)


def test_xlsm_macro_project_rejected(run):
    evil = with_parts(xlsx(), extra={"xl/vbaProject.bin": b"\xd0\xcf\x11\xe0 macro"})
    assert "macro" in text(up(run, "excel-json", evil, "m.xlsx")).lower()


def test_macro_enabled_content_type_rejected(run):
    base = zipfile.ZipFile(io.BytesIO(xlsx())).read("[Content_Types].xml")
    evil = with_parts(xlsx(), replace={"[Content_Types].xml": base.replace(
        b"spreadsheetml.sheet.main+xml", b"ms-excel.sheet.macroEnabled.main+xml")})
    assert "macro-enabled" in text(up(run, "excel-json", evil, "m.xlsx"))


@pytest.mark.parametrize("part,why", [
    ("../../etc/cron.d/evil.xml", "path traversal"),
    ("/abs/evil.xml", "path traversal"),
    ("xl/embeddings/oleObject1.bin", "embedded objects"),
    ("xl/externalLinks/externalLink1.xml", "external workbook links"),
    ("xl/activeX/activeX1.xml", "ActiveX"),
    ("xl/macrosheets/sheet1.xml", "macro sheets"),
    ("customUI/customUI.xml", "custom UI"),
    ("xl/evil.exe", "unexpected file type"),
    ("xl/payload.js", "unexpected file type"),
    ("xl/blob.bin", "opaque binary"),
])
def test_dangerous_parts_rejected(run, part, why):
    evil = with_parts(xlsx(), extra={part: b"<x/>"})
    t = text(up(run, "excel-json", evil, "e.xlsx"))
    assert "Rejected" in t and why in t, (part, t)


def test_xml_dtd_entity_rejected(run):
    xxe = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a SYSTEM "file:///etc/passwd">]><x>&a;</x>'
    evil = with_parts(xlsx(), extra={"xl/worksheets/sheet9.xml": xxe})
    assert "DTD/entity" in text(up(run, "excel-json", evil, "x.xlsx"))
    bomb = b'<!DOCTYPE l [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;&a;">]><l>&b;</l>'
    evil = with_parts(xlsx(), replace={"xl/workbook.xml": bomb})
    assert "DTD/entity" in text(up(run, "excel-json", evil, "x.xlsx"))


def test_zip_bomb_rejected_quickly(run):
    bomb = with_parts(xlsx(), extra={"xl/worksheets/bomb.xml": b"0" * (60 * MB)})   # compresses to ~60 KB
    assert len(bomb) < 1 * MB
    assert "decompression bomb" in text(up(run, "excel-json", bomb, "bomb.xlsx"))


def test_many_entries_and_duplicate_names_rejected(run):
    many = with_parts(xlsx(), extra={f"xl/media/i{n}.png": b"x" for n in range(1100)})
    assert "unusual archive" in text(up(run, "excel-json", many, "many.xlsx"))


def test_encrypted_flag_rejected(run):
    data = bytearray(with_parts(xlsx()))
    # set the "encrypted" general-purpose flag on the first local + central headers
    data[6] |= 1
    idx = data.find(b"PK\x01\x02")
    data[idx + 8] |= 1
    assert "Rejected" in text(up(run, "excel-json", bytes(data), "enc.xlsx"))


def test_screen_unit_never_executes_anything(tmp_path):
    """The screening layer is pure byte inspection: a 'malicious' file in a temp dir is untouched."""
    from app.uploads import screen
    from app.registry import ToolError
    marker = tmp_path / "pwned"
    evil = with_parts(xlsx(), extra={"xl/worksheets/x.xml": f"<x>{marker}</x>".encode(), "evil.sh": b"touch " + str(marker).encode()})
    with pytest.raises(ToolError):
        screen("xlsx", evil)
    assert not marker.exists()


# ------------------------------------------------------------------ images (Base64 image + QR)
def test_image_with_hidden_payload_rejected(run):
    poly = png() + b"MZ\x90\x00 malicious payload appended after IEND"
    for slug, action in (("qr-code", "decode"), ("base64-image", "encode")):
        t = text(up(run, slug, poly, "cat.png", action=action))
        assert "Rejected" in t and "extra data" in t, slug


def test_exe_renamed_png_rejected(run):
    for slug, action in (("qr-code", "decode"), ("base64-image", "encode")):
        t = text(up(run, slug, b"MZ\x90\x00" + b"\0" * 200, "photo.png", action=action))
        assert "Rejected" in t and "Windows executable" in t


def test_svg_with_script_rejected(run):
    t = text(up(run, "base64-image", b'<svg xmlns="http://www.w3.org/2000/svg" onload="alert(1)"/>', "x.png", action="encode"))
    assert "Rejected" in t


def test_truncated_png_rejected(run):
    assert "Rejected" in text(up(run, "qr-code", png()[:-20], "t.png", action="decode"))


def test_genuine_images_still_work(run):
    buf = io.BytesIO(); segno.make("https://example.com").save(buf, kind="png", scale=6)
    r = up(run, "qr-code", buf.getvalue(), "q.png", action="decode")
    assert "https://example.com" in text(r) and "Security check passed" in text(r)
    r = up(run, "base64-image", png(), "ok.png", action="encode")
    assert "Security check passed" in text(r)


def test_pasted_base64_payload_image_also_checked(run):
    poly = base64.b64encode(png() + b"<?php evil ?>").decode()
    assert "Rejected" in text(run("base64-image", input=poly, action="decode"))


# ------------------------------------------------------------------ certificates
def _pem() -> bytes:
    from datetime import timedelta, timezone
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import ec
    from cryptography.x509.oid import NameOID
    key = ec.generate_private_key(ec.SECP256R1())
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "up.example")])
    now = datetime.now(timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key())
            .serial_number(7).not_valid_before(now).not_valid_after(now + timedelta(days=5)).sign(key, hashes.SHA256()))
    return cert.public_bytes(serialization.Encoding.PEM), cert.public_bytes(serialization.Encoding.DER), key.private_bytes(
        serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())


def test_certificate_pem_and_der_uploads_work(run):
    pem, der, _ = _pem()
    assert "CN=up.example" in text(up(run, "certificate", pem, "c.pem", action="decode"))
    r = up(run, "certificate", der, "c.der", action="decode")
    assert "CN=up.example" in text(r) and "Security check passed" in text(r)


def test_private_keys_are_refused_upload_and_paste(run):
    _, _, key = _pem()
    assert "private key" in text(up(run, "certificate", key, "k.pem", action="decode")).lower()
    assert "private key" in text(run("certificate", input=key.decode(), action="decode")).lower()
    pem, _, _ = _pem()
    assert "private key" in text(up(run, "certificate", pem + key, "both.pem", action="decode")).lower()


def test_binary_junk_and_exe_refused_as_certificate(run):
    assert "Rejected" in text(up(run, "certificate", b"MZ" + b"\0" * 100, "c.cer", action="decode"))
    assert "Rejected" in text(up(run, "certificate", b"hello world", "c.pem", action="decode"))
    assert "binary data" in text(up(run, "certificate", b"-----BEGIN CERTIFICATE-----\n\x00\x01", "c.pem", action="decode"))


# ------------------------------------------------------------------ rate limiting / misc
def test_excel_has_stricter_rate_limit():
    from app.registry import get_tool
    assert get_tool("excel-json").rate == "10 per minute"
    assert get_tool("json-format").rate == "30 per minute"


# ------------------------------------------------------------------ capacity protections
def test_excel_parallelism_is_bounded(run):
    from app.tools import excel_json
    held = []
    while excel_json._SLOTS.acquire(blocking=False):       # simulate every conversion slot being busy
        held.append(1)
    try:
        assert "busy" in text(up(run, "excel-json", xlsx(), "a.xlsx", header="on"))
    finally:
        for _ in held:
            excel_json._SLOTS.release()
    assert "busy" not in text(up(run, "excel-json", xlsx(), "a.xlsx", header="on"))   # recovers


def test_healthz(client):
    r = client.get("/healthz")
    assert r.status_code == 200 and r.get_data(as_text=True) == "ok" and r.headers["Cache-Control"] == "no-store"


def test_proxy_hops_use_real_client_ip(monkeypatch):
    from app import create_app
    monkeypatch.setenv("PROXY_HOPS", "1")
    from importlib import reload
    import app.config as cfg
    reload(cfg)
    try:
        app = create_app("testing")
        from flask import request
        app.add_url_rule("/_ip", "ip", lambda: request.remote_addr)
        r = app.test_client().get("/_ip", headers={"X-Forwarded-For": "203.0.113.9, 198.51.100.1"},
                                  environ_base={"REMOTE_ADDR": "10.0.0.1"})
        assert r.get_data(as_text=True) == "198.51.100.1"     # only the last (trusted) hop is believed
    finally:
        monkeypatch.delenv("PROXY_HOPS")
        reload(cfg)
