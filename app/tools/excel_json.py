"""Excel (.xlsx) to JSON.

Safety model: the upload has already passed app.uploads.validate_xlsx (no macros, objects, external
links, DTDs, zip bombs). Here we additionally read the workbook in streaming read-only mode, take
cached VALUES only (formulas are never evaluated), and enforce row / cell / time budgets.
"""
import io
import json
import math
import os
import threading
import time
from datetime import date, datetime, time as dtime, timedelta
from decimal import Decimal

from flask import current_app
from openpyxl import load_workbook

from ..highlight import json_html
from ..registry import Action, Field, Result, Tool, ToolError, register

MAX_SHEETS = 30
MAX_ROWS_PER_SHEET = 20_000
MAX_COLUMNS = 200
MAX_CELLS = 150_000
TIME_BUDGET_S = 10.0
# A dense 1 MB workbook can cost ~2-3 s of CPU. Allow only a few conversions at once per process so a
# burst of uploads cannot starve every other tool; extra requests get a polite "busy" message.
_SLOTS = threading.BoundedSemaphore(max(1, int(os.environ.get("EXCEL_MAX_PARALLEL", "2"))))


def _cell(v):
    """Convert an Excel value to a JSON-safe value."""
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return v if math.isfinite(v) else str(v)
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat()
    if isinstance(v, timedelta):
        return v.total_seconds()
    return str(v)


def _headers(first_row, width: int) -> list[str]:
    seen, out = {}, []
    for i in range(width):
        raw = first_row[i] if i < len(first_row) else None
        name = str(_cell(raw)).strip() if raw not in (None, "") else ""
        name = name or f"column_{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


def _trim(row) -> list:
    row = list(row[:MAX_COLUMNS])
    while row and row[-1] in (None, ""):
        row.pop()
    return row


def convert(data: bytes, sheet: str = "", header: bool = True, include_hidden: bool = False):
    """Return ({sheet_name: rows}, stats). rows are list[dict] (header) or list[list]."""
    started = time.monotonic()
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except Exception as e:  # openpyxl raises many types for malformed input
        current_app.logger.warning("xlsx open failed: %s", type(e).__name__)
        raise ToolError("Could not read the workbook. It may be corrupted.")
    try:
        sheets = [ws for ws in wb.worksheets if include_hidden or ws.sheet_state == "visible"]
        if len(sheets) > MAX_SHEETS:
            raise ToolError(f"Too many sheets (max {MAX_SHEETS}).")
        if sheet:
            by_name = [ws for ws in sheets if ws.title == sheet]
            if not by_name and sheet.isdigit() and 1 <= int(sheet) <= len(sheets):
                by_name = [sheets[int(sheet) - 1]]
            if not by_name:
                raise ToolError("Sheet not found. Use the exact sheet name or its number (1 = first).")
            sheets = by_name
        if not sheets:
            raise ToolError("The workbook has no visible sheets.")

        result, cells, total_rows = {}, 0, 0
        for ws in sheets:
            rows = []
            head = None
            for raw in ws.iter_rows(values_only=True):
                if time.monotonic() - started > TIME_BUDGET_S:
                    raise ToolError("The workbook took too long to process. Try a smaller file.")
                row = _trim(raw)
                if not row:
                    continue                                   # skip empty rows
                if header and head is None:
                    head = _headers(row, len(row))
                    continue
                if len(rows) >= MAX_ROWS_PER_SHEET:
                    raise ToolError(f"Sheet '{ws.title}' has more than {MAX_ROWS_PER_SHEET:,} rows.")
                cells += len(row)
                if cells > MAX_CELLS:
                    raise ToolError(f"The workbook has more than {MAX_CELLS:,} filled cells.")
                vals = [_cell(v) for v in row]
                if header:
                    if len(vals) > len(head):                  # data wider than header row
                        head += _headers([], len(vals))[len(head):]
                    rows.append({head[i]: v for i, v in enumerate(vals)})
                else:
                    rows.append(vals)
            total_rows += len(rows)
            result[ws.title] = rows
        return result, {"sheets": len(result), "rows": total_rows, "cells": cells}
    finally:
        wb.close()


def handler(action, form, files):
    data = files.get("file")
    if not data:
        raise ToolError("Choose an .xlsx file to convert.")
    notes = []
    sheet = form.get("sheet", "").strip()
    if not _SLOTS.acquire(blocking=False):
        raise ToolError("The converter is busy with other files. Please try again in a few seconds.")
    try:
        sheets, stats = convert(data, sheet, form.get("header") == "on", form.get("hidden") == "on")
    finally:
        _SLOTS.release()
    payload = next(iter(sheets.values())) if sheet else sheets       # one sheet -> array, else {name: rows}
    indent = {"2": 2, "4": 4, "compact": None}.get(form.get("indent", "2"), 2)
    try:
        out = json.dumps(payload, ensure_ascii=False, indent=indent,
                         separators=(",", ":") if indent is None else None, allow_nan=False)
    except (ValueError, RecursionError):
        raise ToolError("The workbook contains values that cannot be represented as JSON.")
    notes.append(f"{stats['sheets']} sheet(s), {stats['rows']:,} row(s), {stats['cells']:,} cell(s) converted.")
    return Result(out, html=json_html(out), notes=notes)


register(Tool(
    slug="excel-json", title="Excel to JSON",
    description="Convert an .xlsx workbook to JSON. Macros, links and embedded objects are rejected.",
    fields=(
        Field("file", "Excel file (.xlsx)", "file", accept=".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
              upload="xlsx"),
        Field("header", "First row contains column names", "checkbox", default="on"),
        Field("sheet", "Sheet (name or number; empty = all sheets)", "text", max_len=64),
        Field("indent", "JSON format", "select", (("2", "Pretty (2 spaces)"), ("4", "Pretty (4 spaces)"), ("compact", "Compact")), "2"),
        Field("hidden", "Include hidden sheets", "checkbox"),
    ),
    actions=(Action("convert", "Convert to JSON", True),),
    handler=handler, group="Converters", ext="json", rate="10 per minute",
))
